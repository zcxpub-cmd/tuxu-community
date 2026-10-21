"""Synthetic documents test export integrity, source anchors and hostile text."""

import hashlib
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from PIL import Image
import pytest

from backend.app.documents import MAX_FILE_BYTES, export_document, parse_document, parse_text


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def raster(path: Path, color="steelblue", size=(120, 80)) -> Path:
    Image.new("RGB", size, color).save(path, format="PNG")
    return path


@pytest.fixture
def source_docx(tmp_path):
    existing = raster(tmp_path / "existing.png", "green", (60, 40))
    doc = Document()
    doc.add_heading("自制科学说明样稿", level=1)
    style = doc.styles.add_style("SourceCustomStyle", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Courier New"
    style.font.size = Pt(13)
    doc.add_paragraph("第一段说明装置作用。", style="SourceCustomStyle")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["变量", "数值"], ["温度", "25"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    para = doc.add_paragraph("原稿中的已有图片。")
    para.add_run().add_picture(str(existing), width=Inches(1))
    doc.add_paragraph("第二段说明三个传感单元。")
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


def selection(document, image, text=None, slot_id="slot-a", alt="三个传感单元的概念图"):
    paragraph = next(p for p in document["paragraphs"] if p["kind"] == "paragraph" and p["text"] == (text or "第二段说明三个传感单元。"))
    slot = {"id": slot_id, "paragraph_id": paragraph["id"], "caption": "三个传感单元"}
    chosen = {"slot_id": slot_id, "paragraph_id": paragraph["id"], "image_path": image,
              "caption": "三个传感单元", "alt_text": alt, "asset_id": "asset-a"}
    return slot, chosen


def canonical(element):
    return etree.tostring(element, method="c14n")


def test_docx_root_indices_and_stable_ids(source_docx):
    first, second = parse_document(source_docx), parse_document(source_docx)
    assert [p["id"] for p in first["paragraphs"]] == [p["id"] for p in second["paragraphs"]]
    assert first["paragraphs"][0]["kind"] == "heading"
    assert first["paragraphs"][0]["level"] == 1
    assert next(p for p in first["paragraphs"] if p["kind"] == "table")["source_index"] == 2
    last = next(p for p in first["paragraphs"] if p["text"] == "第二段说明三个传感单元。")
    assert last["source_index"] == 4  # Table must count as one root body block.
    assert first["source_hash"] == digest(source_docx)
    assert next(p for p in first["paragraphs"] if "已有图片" in p["text"])["existing_images"]


def test_docx_preserves_original_and_inserts_at_correct_block(source_docx, tmp_path):
    source_hash = digest(source_docx)
    original = Document(source_docx)
    document = parse_document(source_docx)
    picture = raster(tmp_path / "generated.png")
    slot, chosen = selection(document, picture)
    target = export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
    assert digest(source_docx) == source_hash
    result = Document(target)
    blocks = list(result.element.body)
    assert ParagraphText(blocks[4]) == "第二段说明三个传感单元。"
    assert blocks[5].find(".//" + qn("w:drawing")) is not None
    assert ParagraphText(blocks[6]) == "图 1 三个传感单元（AI 生成示意图）"
    # Existing body blocks, table, custom style and original picture remain intact.
    for old, new in zip(list(original.element.body)[:5], blocks[:5]):
        assert canonical(old) == canonical(new)
    with zipfile.ZipFile(source_docx) as before, zipfile.ZipFile(target) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        assert before.read("word/media/image1.png") == after.read("word/media/image1.png")
    shape = result.inline_shapes[-1]
    assert shape._inline.docPr.get("descr") == chosen["alt_text"]
    assert abs(shape.width / shape.height - 1.5) < 0.00001
    assert blocks[5].find(".//" + qn("w:keepNext")) is not None
    assert blocks[6].find(".//" + qn("w:keepLines")) is not None
    fonts = blocks[6].find(".//" + qn("w:rFonts"))
    assert fonts.get(qn("w:eastAsia")) == "宋体"
    assert fonts.get(qn("w:ascii")) == "Arial"
    # A repeat starts from source, with exactly one additional inline picture.
    second = export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
    repeated = Document(second)
    assert len(repeated.inline_shapes) == len(original.inline_shapes) + 1
    assert len(repeated.paragraphs) == len(original.paragraphs) + 2
    assert digest(source_docx) == source_hash


@pytest.mark.parametrize("pixels", [(768, 1152), (1920, 1080), (1200, 1200)])
@pytest.mark.parametrize("paper", [(8.5, 11, 1.25, 1.0), (6, 8, 0.75, 0.75), (14, 18, 1, 1)])
def test_docx_picture_fits_body_fraction_and_preserves_ratio(tmp_path, pixels, paper):
    """The 768x1152 portrait must be ~3.2x4.8in, never page-wide 6x9in."""
    source = tmp_path / "paper.docx"
    doc = Document()
    section = doc.sections[0]
    page_width, page_height, horizontal_margin, vertical_margin = paper
    section.page_width, section.page_height = Inches(page_width), Inches(page_height)
    section.left_margin = section.right_margin = Inches(horizontal_margin)
    section.top_margin = section.bottom_margin = Inches(vertical_margin)
    doc.add_heading("固定测试稿", 1)
    doc.add_paragraph("第二段说明三个传感单元。")
    doc.add_paragraph("原稿后文保持其文字、顺序和格式。")
    doc.save(source)
    original_hash = digest(source)
    original_blocks = [canonical(block) for block in Document(source).element.body]
    document = parse_document(source)
    slot, chosen = selection(document, raster(tmp_path / "picture.png", size=pixels))
    output = export_document(source, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
    exported = Document(output)
    shape = exported.inline_shapes[0]
    actual_section = exported.sections[0]
    max_width = min(actual_section.page_width - actual_section.left_margin - actual_section.right_margin, Inches(6))
    max_height = min(int((actual_section.page_height - actual_section.top_margin - actual_section.bottom_margin) * 0.55), Inches(4.8))
    scale = min(max_width / pixels[0], max_height / pixels[1])
    assert abs(shape.width - pixels[0] * scale) <= 1
    assert abs(shape.height - pixels[1] * scale) <= 1
    assert shape.width <= max_width and shape.height <= max_height
    assert abs(shape.width / shape.height - pixels[0] / pixels[1]) < 0.000001
    if pixels == (1920, 1080):
        assert abs(shape.width - max_width) <= 1, "Ordinary landscape keeps its existing maximum width"
    if pixels == (768, 1152) and paper[1] == 18:
        assert abs(shape.height - Inches(4.8)) <= 1, "Large paper still caps paragraph illustrations at 4.8in"
    blocks = list(exported.element.body)
    assert blocks[2].find(".//" + qn("w:drawing")) is not None
    assert blocks[2].find(".//" + qn("w:keepNext")) is not None
    assert blocks[3].find(".//" + qn("w:keepLines")) is not None
    assert [canonical(block) for index, block in enumerate(blocks) if index not in (2, 3)] == original_blocks
    assert digest(source) == original_hash


def test_docx_picture_uses_section_at_insertion_without_moving_break(tmp_path):
    source = tmp_path / "sections.docx"
    doc = Document()
    first = doc.sections[0]
    first.page_width, first.page_height = Inches(6), Inches(8)
    first.left_margin = first.right_margin = first.top_margin = first.bottom_margin = Inches(0.5)
    doc.add_paragraph("第一节正文锚点。")
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    # This paragraph retains the original section break. Inserting after it
    # places the image in section 2, not section 1, without relocating the break.
    doc.paragraphs[-1].add_run("第一节末尾正文锚点。")
    second.page_width, second.page_height = Inches(8.5), Inches(11)
    second.left_margin = second.right_margin = second.top_margin = second.bottom_margin = Inches(1)
    doc.add_paragraph("第二节正文锚点。")
    doc.add_paragraph("第二节后文。")
    doc.save(source)
    source_hash = digest(source)
    source_blocks = [canonical(block) for block in Document(source).element.body]
    document = parse_document(source)
    image = raster(tmp_path / "portrait.png", size=(768, 1152))
    items = [selection(document, image, text, slot_id=slot_id) for slot_id, text in
             [("first", "第一节正文锚点。"), ("boundary", "第一节末尾正文锚点。"), ("second", "第二节正文锚点。")]]
    output = export_document(source, document, {"slots": [item[0] for item in reversed(items)]},
                             [item[1] for item in reversed(items)], tmp_path / "out", "docx")
    result = Document(output)
    shapes = list(result.inline_shapes)
    assert len(shapes) == 3
    assert abs(shapes[0].height - Inches(7) * 0.55) <= 1
    assert abs(shapes[1].height - Inches(4.8)) <= 1
    assert abs(shapes[2].height - Inches(4.8)) <= 1
    assert all(abs(shape.width / shape.height - 2 / 3) < 0.000001 for shape in shapes)
    blocks = list(result.element.body)
    retained = []
    skip_caption = False
    for block in blocks:
        if block.find(".//" + qn("w:drawing")) is not None:
            skip_caption = True
            continue
        if skip_caption:
            assert ParagraphText(block).startswith("图 ")
            skip_caption = False
            continue
        retained.append(canonical(block))
    assert retained == source_blocks, "Original text and section breaks must not move or change"
    assert digest(source) == source_hash


def ParagraphText(element):
    return "".join(element.itertext()) if element.tag != qn("w:p") else "".join(node.text or "" for node in element.iter(qn("w:t")))


def test_numbering_follows_source_position_not_selection_order(source_docx, tmp_path):
    document = parse_document(source_docx)
    image = raster(tmp_path / "new.png")
    late_slot, late = selection(document, image, slot_id="late")
    early_slot, early = selection(document, image, "第一段说明装置作用。", slot_id="early")
    target = export_document(source_docx, document, {"slots": [late_slot, early_slot]}, [late, early], tmp_path / "out", "docx")
    captions = [p.text for p in Document(target).paragraphs if p.text.startswith("图 ")]
    assert captions == ["图 1 三个传感单元（AI 生成示意图）", "图 2 三个传感单元（AI 生成示意图）"]
    blocks = list(Document(target).element.body)
    assert ParagraphText(blocks[1]) == "第一段说明装置作用。"
    assert ParagraphText(blocks[3]).startswith("图 1 ")
    assert ParagraphText(blocks[6]) == "第二段说明三个传感单元。"


@pytest.mark.parametrize("mutation", ["revision", "columns", "field", "nested", "floating"])
def test_rejects_unsupported_docx_structures(tmp_path, mutation):
    doc = Document()
    para = doc.add_paragraph("可阅读正文")
    if mutation == "revision":
        para._p.append(OxmlElement("w:ins"))
    elif mutation == "columns":
        doc.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    elif mutation == "field":
        para.add_run()._r.append(OxmlElement("w:instrText"))
    elif mutation == "nested":
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).add_table(rows=1, cols=1)
    else:
        drawing = OxmlElement("w:drawing")
        drawing.append(OxmlElement("wp:anchor"))
        para.add_run()._r.append(drawing)
    path = tmp_path / f"{mutation}.docx"
    doc.save(path)
    with pytest.raises(ValueError):
        parse_document(path)


def test_limit_and_invalid_text_files(tmp_path):
    assert parse_text("字" * 10_000, "限制内")["character_count"] == 10_000
    with pytest.raises(ValueError, match="10000"):
        parse_text("字" * 10_001)
    with pytest.raises(ValueError, match="正文"):
        parse_text(" \n\t")
    with pytest.raises(ValueError, match="控制字符"):
        parse_text("正文\x00注入")
    oversized = tmp_path / "large.txt"
    with oversized.open("wb") as handle:
        handle.truncate(MAX_FILE_BYTES + 1)
    with pytest.raises(ValueError, match="10 MB"):
        parse_document(oversized)
    invalid = tmp_path / "invalid.txt"
    invalid.write_bytes(b"\xff\xfe")
    with pytest.raises(ValueError, match="UTF-8"):
        parse_document(invalid)


def test_docx_rejects_external_template_but_keeps_normal_hyperlink(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    doc.add_paragraph("正文中允许保留普通超链接。")
    doc.part.relate_to("https://example.invalid/reference", RT.HYPERLINK, is_external=True)
    safe = tmp_path / "hyperlink.docx"
    doc.save(safe)
    assert parse_document(safe)["character_count"] > 0
    template_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate"
    doc.part.relate_to("https://example.invalid/template.dotx", template_type, is_external=True)
    unsafe = tmp_path / "template.docx"
    doc.save(unsafe)
    with pytest.raises(ValueError, match="外部链接资源"):
        parse_document(unsafe)


@pytest.mark.parametrize("kind", ["heading", "table"])
def test_export_rejects_non_body_and_unknown_anchors(source_docx, tmp_path, kind):
    document = parse_document(source_docx)
    paragraph = next(p for p in document["paragraphs"] if p["kind"] == kind)
    image = raster(tmp_path / "new.png")
    slot = {"id": "bad", "paragraph_id": paragraph["id"]}
    chosen = {"slot_id": "bad", "paragraph_id": paragraph["id"], "image_path": image}
    with pytest.raises(ValueError, match="普通正文"):
        export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
    chosen["paragraph_id"] = "missing"
    with pytest.raises(ValueError, match="不匹配"):
        export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")


def test_source_change_collision_and_tampered_anchor_rejected(source_docx, tmp_path):
    document = parse_document(source_docx)
    image = raster(tmp_path / "new.png")
    slot, chosen = selection(document, image)
    original_index = next(p for p in document["paragraphs"] if p["id"] == chosen["paragraph_id"])
    original_index["source_index"] = 1
    with pytest.raises(ValueError, match="锚点"):
        export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
    document = parse_document(source_docx)
    collision = source_docx.parent / "illustrated.docx"
    collision.write_bytes(source_docx.read_bytes())
    with pytest.raises(ValueError, match="覆盖原稿"):
        export_document(collision, document, {"slots": [slot]}, [chosen], collision.parent, "docx")
    source_docx.write_bytes(source_docx.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="原稿已改变"):
        export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")


def test_markdown_source_preservation_and_relative_bundle(tmp_path):
    text = '# 标题\n\n正文含 [原始链接](https://example.invalid)。\n继续这一段。\n\n```python\nprint("保留代码")\n```\n\n| 变量 | 值 |\n| --- | --- |\n| A | 2 |\n'
    source = tmp_path / "article.md"
    source.write_text(text, encoding="utf-8")
    document = parse_document(source)
    image = raster(tmp_path / "new.png")
    ordinary = next(p for p in document["paragraphs"] if p["insertable"])
    slot = {"id": "new", "paragraph_id": ordinary["id"]}
    chosen = {"slot_id": "new", "paragraph_id": ordinary["id"], "image_path": image, "caption": "概念说明", "alt_text": "说明"}
    target = export_document(source, document, {"slots": [slot]}, [chosen], tmp_path / "out", "markdown")
    with zipfile.ZipFile(target) as package:
        assert set(package.namelist()) == {"article.md", "images/figure_001.png"}
        exported = package.read("article.md").decode("utf-8")
        assert "[原始链接](https://example.invalid)" in exported
        assert '```python\nprint("保留代码")\n```' in exported
        assert "| A | 2 |" in exported
        assert "继续这一段。\n\n![说明](images/figure_001.png)" in exported
        assert "AI 生成示意图" in exported
    assert source.read_text(encoding="utf-8") == text
    code = next(p for p in document["paragraphs"] if p.get("protected"))
    slot["paragraph_id"] = chosen["paragraph_id"] = code["id"]
    with pytest.raises(ValueError, match="代码块"):
        export_document(source, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")


def test_html_self_contained_and_escapes_source_and_attributes(tmp_path):
    text = '<script>alert("source")</script> & 正文'
    document = parse_text(text, '<img src=x onerror="alert(1)">')
    image = raster(tmp_path / "new.png")
    slot, chosen = selection(document, image, text=text, alt='" onerror="alert(2) <script>')
    chosen["caption"] = '<script>alert("caption")</script>'
    target = export_document(None, document, {"slots": [slot]}, [chosen], tmp_path / "out", "html")
    exported = target.read_text(encoding="utf-8")
    assert "<script>" not in exported
    assert '<img src=x' not in exported
    assert "&lt;script&gt;" in exported
    assert 'src="data:image/png;base64,' in exported
    assert 'alt="&quot; onerror=&quot;alert(2) &lt;script&gt;"' in exported
    assert "http://" not in exported and "file://" not in exported
    parsed = etree.HTML(exported)
    assert len(parsed.xpath("//script")) == 0
    assert len(parsed.xpath("//img[@onerror]")) == 0
    assert "Content-Security-Policy" in exported


def test_txt_to_markdown_escapes_syntax_and_supports_no_image_plan(tmp_path):
    document = parse_text("# 普通文本\n<script>alert(1)</script>\n**不是Markdown强调**", "TXT稿")
    target = export_document(None, document, {"slots": []}, [], tmp_path / "out", "markdown")
    with zipfile.ZipFile(target) as package:
        markdown = package.read("article.md").decode("utf-8")
        assert "\\# 普通文本" in markdown
        assert "\\<script\\>" in markdown
        assert "\\*\\*不是Markdown强调\\*\\*" in markdown
        assert set(package.namelist()) == {"article.md"}


def test_docx_existing_images_and_table_survive_preview_and_bundle(source_docx, tmp_path):
    document = parse_document(source_docx)
    image = raster(tmp_path / "new.png")
    slot, chosen = selection(document, image)
    preview = export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "html")
    output = preview.read_text(encoding="utf-8")
    assert output.count('src="data:image/png;base64,') == 2
    assert "<td>温度</td>" in output and "<td>25</td>" in output
    bundle = export_document(source_docx, document, {"slots": [slot]}, [chosen], tmp_path / "out", "markdown")
    with zipfile.ZipFile(bundle) as package:
        assert "images/source_0003_01.png" in package.namelist()
        assert "![原稿图片](images/source_0003_01.png)" in package.read("article.md").decode("utf-8")


def test_duplicate_selection_and_missing_docx_source_rejected(source_docx, tmp_path):
    document = parse_document(source_docx)
    image = raster(tmp_path / "new.png")
    slot, chosen = selection(document, image)
    with pytest.raises(ValueError, match="只能选择一张"):
        export_document(source_docx, document, {"slots": [slot]}, [chosen, chosen], tmp_path / "out", "docx")
    with pytest.raises(ValueError, match="原始文件"):
        export_document(None, document, {"slots": [slot]}, [chosen], tmp_path / "out", "docx")
