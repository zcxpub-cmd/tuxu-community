"""Bounded document parsing and reversible illustration exports.

The source document is never modified. DOCX paragraph source_index values point
at *root body blocks*, including tables and image-only paragraphs.
"""

from __future__ import annotations

import base64
import hashlib
import html
import io
import os
from pathlib import Path, PurePosixPath
import re
import tempfile
from typing import Any
from uuid import uuid4
import zipfile

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image, UnidentifiedImageError


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARACTERS = 10_000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_IMAGE_BYTES = 30 * 1024 * 1024
MAX_IMAGE_PIXELS = 50_000_000
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_HEADING = re.compile(r"^ {0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
_FENCE = re.compile(r"^ {0,3}(`{3,}|~{3,})")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _check_text(text: str) -> None:
    if _ILLEGAL_XML.search(text):
        raise ValueError("稿件包含不支持的控制字符，请先清理后导入")


def _paragraph(block_index: int, text: str, kind: str = "paragraph", level: int = 0, **extra: Any) -> dict:
    _check_text(text)
    digest = _hash(f"{kind}\0{text}".encode("utf-8"))[:12]
    return {"id": f"p_{block_index}_{digest}", "text": text, "kind": kind,
            "level": level, "source_index": block_index,
            "insertable": kind == "paragraph" and bool(text.strip()), **extra}


def _finish(paragraphs: list[dict], *, title: str, format: str, source_hash: str,
            source_filename: str, **extra: Any) -> dict:
    if not any(p["text"].strip() for p in paragraphs):
        raise ValueError("稿件中没有可阅读的正文")
    character_count = sum(len(p["text"]) for p in paragraphs)
    if character_count > MAX_TEXT_CHARACTERS:
        raise ValueError("正文超过 10000 字符，首版暂不支持")
    heading = next((p["text"] for p in paragraphs if p["kind"] == "heading"), None)
    return {"id": str(uuid4()), "title": title or heading or Path(source_filename).stem or "未命名稿件",
            "format": format, "source_hash": source_hash, "paragraphs": paragraphs,
            "warnings": [], "character_count": character_count,
            "source_filename": source_filename, **extra}


def _table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _parse_text_content(text: str, title: str, format: str, source_hash: str, source_filename: str) -> dict:
    _check_text(text)
    lines = text.splitlines(keepends=True)
    paragraphs: list[dict] = []
    index = 0
    while index < len(lines):
        if not lines[index].strip():
            index += 1
            continue
        start = index
        value = lines[index].rstrip("\r\n")
        heading = _HEADING.match(value) if format == "md" else None
        fence = _FENCE.match(value) if format == "md" else None
        if heading:
            paragraphs.append(_paragraph(start, heading.group(2), "heading", len(heading.group(1)),
                                         source_line_start=start, source_line_end=start))
            index += 1
        elif fence:
            marker = fence.group(1)[0]
            length = len(fence.group(1))
            index += 1
            while index < len(lines):
                closing = re.match(rf"^ {{0,3}}{re.escape(marker)}{{{length},}}\s*$", lines[index].rstrip("\r\n"))
                index += 1
                if closing:
                    break
            paragraphs.append(_paragraph(start, "".join(lines[start:index]).rstrip("\r\n"),
                                         insertable=False, protected=True,
                                         source_line_start=start, source_line_end=index - 1))
        elif format == "md" and index + 1 < len(lines) and _TABLE_RULE.match(lines[index + 1].rstrip("\r\n")):
            rows = [_table_cells(value)]
            index += 2
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                rows.append(_table_cells(lines[index]))
                index += 1
            paragraphs.append(_paragraph(start, "\n".join(" | ".join(row) for row in rows), "table",
                                         rows=rows, source_line_start=start, source_line_end=index - 1))
        else:
            index += 1
            # TXT has explicit line paragraphs. MD wraps soft lines into one block.
            if format == "md":
                while index < len(lines) and lines[index].strip():
                    next_line = lines[index].rstrip("\r\n")
                    if _HEADING.match(next_line) or _FENCE.match(next_line):
                        break
                    if index + 1 < len(lines) and _TABLE_RULE.match(lines[index + 1].rstrip("\r\n")):
                        break
                    index += 1
            paragraphs.append(_paragraph(start, "".join(lines[start:index]).rstrip("\r\n"),
                                         source_line_start=start, source_line_end=index - 1))
    return _finish(paragraphs, title=title, format=format, source_hash=source_hash,
                   source_filename=source_filename, source_text=text)


def parse_text(text: str, title: str = "") -> dict:
    if not isinstance(text, str):
        raise ValueError("正文必须是文本")
    data = text.encode("utf-8")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("文件超过 10 MB，首版暂不支持")
    return _parse_text_content(text, title, "txt", _hash(data), "pasted.txt")


def _validate_docx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            names = package.namelist()
            if len(names) != len(set(names)) or any(
                name.startswith("/") or ".." in PurePosixPath(name).parts for name in names
            ):
                raise ValueError("DOCX 包路径异常")
            if sum(info.file_size for info in package.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX 解压后过大，首版暂不支持")
            if "word/document.xml" not in names:
                raise ValueError("文件不是可读取的 DOCX")
            parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
            unsupported = {
                "ins": "修订追踪", "del": "修订追踪", "moveFrom": "修订追踪", "moveTo": "修订追踪",
                "pPrChange": "修订追踪", "rPrChange": "修订追踪", "tblPrChange": "修订追踪",
                "sectPrChange": "修订追踪", "tblGridChange": "修订追踪", "tcPrChange": "修订追踪",
                "trPrChange": "修订追踪", "cellIns": "修订追踪", "cellDel": "修订追踪",
                "cellMerge": "修订追踪", "numberingChange": "修订追踪",
                "fldSimple": "域或交叉引用", "fldChar": "复杂域或交叉引用",
                "instrText": "复杂域或交叉引用", "sdt": "内容控件", "altChunk": "嵌入内容",
                "txbxContent": "文本框", "object": "嵌入对象", "pict": "旧式图片或形状",
                "footnoteReference": "脚注", "endnoteReference": "尾注",
            }
            for name in names:
                if name.endswith(".rels"):
                    relations = etree.fromstring(package.read(name), parser)
                    if any(
                        rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink")
                        for rel in relations
                    ):
                        raise ValueError("DOCX 含外部链接资源，请先内嵌或断开外部资源")
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                xml = package.read(name)
                if b"<!DOCTYPE" in xml.upper():
                    raise ValueError("DOCX 包含不支持的 XML 声明")
                root = etree.fromstring(xml, parser)
                for element in root.iter():
                    if not isinstance(element.tag, str):
                        continue
                    tag = etree.QName(element).localname
                    namespace = etree.QName(element).namespace
                    if namespace == W_NS and tag in unsupported:
                        raise ValueError(f"DOCX 含{unsupported[tag]}，请先转为普通文档")
                    if namespace == W_NS and tag == "trackRevisions" and element.get(qn("w:val"), "1") not in {"0", "false", "off"}:
                        raise ValueError("DOCX 启用了修订追踪，请先关闭并确认修订")
                    if namespace == W_NS and tag == "cols":
                        if int(element.get(qn("w:num"), "1")) > 1 or len(element.findall(qn("w:col"))) > 1:
                            raise ValueError("DOCX 含双栏或多栏，首版暂不支持")
                    if tag == "anchor" and namespace and "wordprocessingDrawing" in namespace:
                        raise ValueError("DOCX 含浮动图片，首版仅支持行内图片")
                    if namespace == W_NS and tag == "tbl" and element.findall(".//w:tbl", NS):
                        raise ValueError("DOCX 含嵌套表格，首版暂不支持")
                    if namespace == W_NS and tag in {"gridSpan", "vMerge"}:
                        raise ValueError("DOCX 含合并单元格，首版仅支持基础表格")
                    if tag == "chart" and namespace and "drawingml" in namespace:
                        raise ValueError("DOCX 含嵌入图表，首版暂不支持")
    except (zipfile.BadZipFile, etree.XMLSyntaxError, KeyError) as exc:
        raise ValueError("DOCX 文件损坏或无法读取") from exc


def _heading_level(paragraph: Paragraph) -> int:
    properties = paragraph._p.pPr
    outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
    if outline is not None and int(outline.get(qn("w:val"), "9")) < 9:
        return int(outline.get(qn("w:val"))) + 1
    style = paragraph.style
    seen: set[str] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        if style.name in {"Title", "标题"}:
            return 1
        if style.name in {"Subtitle", "副标题"}:
            return 2
        match = re.match(r"(?:heading|标题)\s*([1-9])", style.name, re.IGNORECASE)
        if match:
            return int(match.group(1))
        properties = style.element.find(qn("w:pPr"))
        outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
        if outline is not None and int(outline.get(qn("w:val"), "9")) < 9:
            return int(outline.get(qn("w:val"))) + 1
        style = style.base_style
    return 0


def _existing_images(block: Any, doc: Any) -> list[dict]:
    images = []
    for element in block.iter():
        if element.tag != f"{{{NS['a']}}}blip":
            continue
        rid = element.get(qn("r:embed"))
        if not rid or rid not in doc.part.related_parts:
            raise ValueError("DOCX 含外部链接图片，首版需要内嵌图片")
        part = doc.part.related_parts[rid]
        _image_info(part.blob)
        images.append({"relationship_id": rid, "content_type": part.content_type})
    return images


def _parse_docx(data: bytes, filename: str) -> dict:
    _validate_docx(data)
    doc = WordDocument(io.BytesIO(data))
    paragraphs = []
    for source_index, block in enumerate(doc.element.body):
        if block.tag == qn("w:p"):
            para = Paragraph(block, doc)
            text = para.text
            level = _heading_level(para)
            images = _existing_images(block, doc)
            if text.strip() or images:
                paragraphs.append(_paragraph(source_index, text, "heading" if level else "paragraph", level,
                                             existing_images=images))
        elif block.tag == qn("w:tbl"):
            table = Table(block, doc)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            paragraphs.append(_paragraph(source_index, "\n".join(" | ".join(row) for row in rows), "table",
                                         rows=rows, existing_images=_existing_images(block, doc)))
        elif block.tag != qn("w:sectPr"):
            raise ValueError("DOCX 含不支持的正文结构，请先转为普通文档")
    return _finish(paragraphs, title="", format="docx", source_hash=_hash(data), source_filename=filename)


def parse_document(path: Path) -> dict:
    path = Path(path)
    if not path.is_file():
        raise ValueError("源文件不存在")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过 10 MB，首版暂不支持")
    suffix = path.suffix.lower()
    if suffix not in {".docx", ".md", ".txt"}:
        raise ValueError("首版仅支持 DOCX、Markdown 和 TXT")
    data = path.read_bytes()
    if suffix == ".docx":
        return _parse_docx(data, path.name)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("文本文件必须是 UTF-8 编码") from exc
    return _parse_text_content(text, "", suffix[1:], _hash(data), path.name)


def _image_info(data: bytes) -> tuple[str, int, int]:
    if len(data) > MAX_IMAGE_BYTES:
        raise ValueError("图片文件过大")
    try:
        with Image.open(io.BytesIO(data)) as image:
            width, height = image.size
            if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                raise ValueError("图片像素超出安全范围")
            format = image.format
            if format not in {"PNG", "JPEG", "GIF", "BMP", "TIFF", "WEBP"}:
                raise ValueError("图片格式不受支持")
            image.verify()
            return format, width, height
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ValueError("图片损坏或无法解码") from exc


def _caption(number: int, text: str) -> str:
    text = re.sub(r"^\s*图\s*\d+\s*[：:.、-]?\s*", "", str(text or "配图示意"))
    _check_text(text)
    if "AI生成示意图" not in re.sub(r"\s+", "", text):
        text = f"{text}（AI 生成示意图）"
    return f"图 {number} {text}"


def _validated_selections(document: dict, plan: dict, selections: list[dict]) -> list[dict]:
    paragraphs = {p["id"]: p for p in document["paragraphs"]}
    slots = plan.get("slots", [])
    if len(slots) > 5 or len(selections) > 5:
        raise ValueError("首版最多支持 5 处图位")
    slot_map = {slot["id"]: slot for slot in slots}
    if len(slot_map) != len(slots):
        raise ValueError("方案图位 ID 不能重复")
    slot_order = {slot["id"]: index for index, slot in enumerate(slots)}
    seen: set[str] = set()
    result = []
    for selection in selections:
        slot_id = selection.get("slot_id")
        slot = slot_map.get(slot_id)
        paragraph = paragraphs.get(selection.get("paragraph_id"))
        if slot is None or paragraph is None or slot.get("paragraph_id") != paragraph["id"]:
            raise ValueError("选图与当前方案或段落不匹配")
        if paragraph["kind"] != "paragraph" or not paragraph.get("insertable", bool(paragraph["text"].strip())):
            raise ValueError("图位只能位于普通正文段落，不能插入标题、表格或代码块")
        if slot_id in seen:
            raise ValueError("每个图位只能选择一张图片")
        seen.add(slot_id)
        image_path = Path(selection["image_path"])
        if not image_path.is_file() or image_path.stat().st_size > MAX_IMAGE_BYTES:
            raise ValueError("所选图片不存在或过大")
        data = image_path.read_bytes()
        image_format, width, height = _image_info(data)
        alt_text = str(selection.get("alt_text") or selection.get("caption") or "AI 生成示意图")
        _check_text(alt_text)
        result.append({**selection, "paragraph": paragraph, "image_data": data,
                       "image_format": image_format, "width": width, "height": height,
                       "alt_text": alt_text, "slot_order": slot_order[slot_id]})
    result.sort(key=lambda row: (row["paragraph"]["source_index"], row["slot_order"]))
    for index, row in enumerate(result, 1):
        row["caption"] = _caption(index, row.get("caption", slot_map[row["slot_id"]].get("caption", "")))
        row["number"] = index
    return result


def _fonts(run: Any, size: float = 11) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def _picture_bytes(row: dict) -> bytes:
    if row["image_format"] != "WEBP":
        return row["image_data"]
    with Image.open(io.BytesIO(row["image_data"])) as image:
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        return stream.getvalue()


def _section_picture_bounds(doc: Any) -> list[tuple[int, int]]:
    # Missing section properties inherit from the preceding section. Use the
    # Word template only for properties absent from the first section too.
    defaults = WordDocument().sections[0]
    properties = ("page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin")
    dimensions = {name: getattr(defaults, name) for name in properties}
    result = []
    for section in doc.sections:
        for name in properties:
            value = getattr(section, name)
            if value is not None:
                dimensions[name] = value
        content_width = dimensions["page_width"] - dimensions["left_margin"] - dimensions["right_margin"]
        content_height = dimensions["page_height"] - dimensions["top_margin"] - dimensions["bottom_margin"]
        if content_width <= 0 or content_height <= 0:
            raise ValueError("文档页边距没有可用的插图空间")
        # A paragraph illustration should leave room for its caption and nearby
        # prose, rather than expand a portrait image to nearly a whole page.
        result.append((min(content_width, Inches(6)), min(int(content_height * 0.55), Inches(4.8))))
    return result


def _write_docx(source_data: bytes | None, document: dict, rows: list[dict], target: Path) -> None:
    doc = WordDocument(io.BytesIO(source_data)) if source_data is not None else WordDocument()
    anchors = {}
    anchor_bounds = {}
    section_bounds = _section_picture_bounds(doc)
    if source_data is not None:
        blocks = list(doc.element.body)
        bounds_by_block = {}
        section_index = 0
        for index, block in enumerate(blocks):
            # A sectPr on the anchor ends that section; new paragraphs inserted
            # after it belong to the following section without moving the break.
            if block.tag == qn("w:p") and block.find("w:pPr/w:sectPr", NS) is not None:
                section_index += 1
            bounds_by_block[index] = section_bounds[min(section_index, len(section_bounds) - 1)]
        for paragraph in document["paragraphs"]:
            index = paragraph["source_index"]
            if paragraph["kind"] == "paragraph":
                if not isinstance(index, int) or isinstance(index, bool) or index >= len(blocks) or index < 0 or blocks[index].tag != qn("w:p"):
                    raise ValueError("原文段落位置已变化，不能可靠插图")
                anchors[paragraph["id"]] = blocks[index]
                anchor_bounds[paragraph["id"]] = bounds_by_block[index]
    else:
        for paragraph in document["paragraphs"]:
            if paragraph["kind"] == "table":
                values = paragraph.get("rows", [])
                if values:
                    columns = max(len(value) for value in values)
                    table = doc.add_table(rows=len(values), cols=columns)
                    table.style = "Table Grid"
                    for row_index, value in enumerate(values):
                        for column_index, text in enumerate(value):
                            table.cell(row_index, column_index).text = text
                continue
            para = doc.add_paragraph()
            if paragraph["kind"] == "heading":
                para.style = f"Heading {min(paragraph['level'] or 1, 9)}"
            size = 18 - 2 * (min(paragraph["level"] or 1, 4) - 1) if paragraph["kind"] == "heading" else 11
            _fonts(para.add_run(paragraph["text"]), size)
            anchors[paragraph["id"]] = para._p
            anchor_bounds[paragraph["id"]] = section_bounds[0]
    last_insert = {}
    for row in rows:
        anchor_id = row["paragraph"]["id"]
        anchor = last_insert.get(anchor_id, anchors.get(anchor_id))
        if anchor is None:
            raise ValueError("找不到图位对应的正文段落")
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        picture.paragraph_format.keep_together = True
        max_width, max_height = anchor_bounds[anchor_id]
        scale = min(max_width / row["width"], max_height / row["height"])
        shape = picture.add_run().add_picture(io.BytesIO(_picture_bytes(row)),
                                              width=int(row["width"] * scale), height=int(row["height"] * scale))
        shape._inline.docPr.set("descr", row["alt_text"])
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_together = True
        caption.paragraph_format.keep_with_next = False
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(8)
        _fonts(caption.add_run(row["caption"]), 10.5)
        anchor.addnext(picture._p)
        picture._p.addnext(caption._p)
        last_insert[anchor_id] = caption._p
    doc.save(target)


def _source_images(source_data: bytes | None, document: dict) -> dict[str, list[tuple[bytes, str]]]:
    if source_data is None:
        return {}
    doc = WordDocument(io.BytesIO(source_data))
    result = {}
    for paragraph in document["paragraphs"]:
        images = []
        for metadata in paragraph.get("existing_images", []):
            part = doc.part.related_parts[metadata["relationship_id"]]
            data = part.blob
            format, _, _ = _image_info(data)
            images.append((data, format))
        result[paragraph["id"]] = images
    return result


def _markdown_escape(text: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>#|~])", r"\\\1", text)


def _write_markdown(document: dict, rows: list[dict], source_data: bytes | None, target: Path) -> None:
    by_paragraph: dict[str, list[dict]] = {}
    files: dict[str, bytes] = {}
    for row in rows:
        suffix = {"JPEG": "jpg", "TIFF": "tiff"}.get(row["image_format"], row["image_format"].lower())
        row["bundle_path"] = f"images/figure_{row['number']:03d}.{suffix}"
        files[row["bundle_path"]] = row["image_data"]
        by_paragraph.setdefault(row["paragraph"]["id"], []).append(row)
    existing = _source_images(source_data, document)

    def figures(paragraph: dict) -> str:
        fragments = []
        for index, (data, format) in enumerate(existing.get(paragraph["id"], []), 1):
            suffix = {"JPEG": "jpg", "TIFF": "tiff"}.get(format, format.lower())
            path = f"images/source_{paragraph['source_index']:04d}_{index:02d}.{suffix}"
            files[path] = data
            fragments.append(f"![原稿图片]({path})")
        for row in by_paragraph.get(paragraph["id"], []):
            fragments.append(f"![{_markdown_escape(row['alt_text'])}]({row['bundle_path']})\n\n{_markdown_escape(row['caption'])}")
        return "\n\n".join(fragments)

    if document["format"] == "md" and "source_text" in document:
        lines = document["source_text"].splitlines(keepends=True)
        insertions = {}
        for paragraph in document["paragraphs"]:
            fragment = figures(paragraph)
            if fragment:
                insertions[paragraph["source_line_end"]] = fragment
        parts = []
        for index, line in enumerate(lines):
            parts.append(line)
            if index in insertions:
                if not line.endswith(("\n", "\r")):
                    parts.append("\n")
                parts.append(f"\n{insertions[index]}\n\n")
        markdown = "".join(parts)
    else:
        parts = []
        for paragraph in document["paragraphs"]:
            text = _markdown_escape(paragraph["text"])
            if paragraph["kind"] == "heading":
                text = f"{'#' * min(paragraph['level'] or 1, 6)} {text}"
            elif paragraph["kind"] == "table" and paragraph.get("rows"):
                values = paragraph["rows"]
                columns = max(len(value) for value in values)
                render = lambda value: "| " + " | ".join(_markdown_escape(str(cell)).replace("\n", "<br>") for cell in value + [""] * (columns - len(value))) + " |"
                text = "\n".join([render(values[0]), "| " + " | ".join(["---"] * columns) + " |", *[render(value) for value in values[1:]]])
            if text:
                parts.append(text)
            fragment = figures(paragraph)
            if fragment:
                parts.append(fragment)
        markdown = "\n\n".join(parts) + "\n"
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr("article.md", markdown.encode("utf-8"))
        for path, data in files.items():
            package.writestr(path, data)


def _data_image(data: bytes, format: str, alt: str) -> str:
    mime = {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif", "BMP": "image/bmp", "TIFF": "image/tiff", "WEBP": "image/webp"}[format]
    # Convert formats without dependable browser display; embedded raster only.
    if format in {"BMP", "TIFF"}:
        with Image.open(io.BytesIO(data)) as image:
            stream = io.BytesIO()
            image.save(stream, format="PNG")
            data, mime = stream.getvalue(), "image/png"
    return f'<img src="data:{mime};base64,{base64.b64encode(data).decode("ascii")}" alt="{html.escape(alt, quote=True)}">'


def _write_html(document: dict, rows: list[dict], source_data: bytes | None, target: Path) -> None:
    by_paragraph: dict[str, list[dict]] = {}
    for row in rows:
        by_paragraph.setdefault(row["paragraph"]["id"], []).append(row)
    existing = _source_images(source_data, document)
    parts = []
    for paragraph in document["paragraphs"]:
        text = html.escape(paragraph["text"], quote=True).replace("\n", "<br>")
        if paragraph["kind"] == "table":
            values = paragraph.get("rows", [[paragraph["text"]]])
            cells = "".join("<tr>" + "".join(f"<td>{html.escape(str(cell), quote=True).replace(chr(10), '<br>')}</td>" for cell in value) + "</tr>" for value in values)
            parts.append(f"<table>{cells}</table>")
        else:
            tag = f"h{min(paragraph['level'] or 1, 6)}" if paragraph["kind"] == "heading" else "p"
            parts.append(f"<{tag}>{text}</{tag}>")
        for data, format in existing.get(paragraph["id"], []):
            parts.append(f"<figure>{_data_image(data, format, '原稿图片')}</figure>")
        for row in by_paragraph.get(paragraph["id"], []):
            parts.append(f"<figure>{_data_image(row['image_data'], row['image_format'], row['alt_text'])}<figcaption>{html.escape(row['caption'], quote=True)}</figcaption></figure>")
    output = '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
    output += '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">'
    output += f"<title>{html.escape(document.get('title', '文章预览'), quote=True)}</title>"
    output += '<style>body{font:17px/1.8 Arial,"宋体",serif;max-width:850px;margin:40px auto;padding:0 24px;overflow-wrap:anywhere}figure{text-align:center;margin:24px 0;break-inside:avoid}img{max-width:100%;height:auto}figcaption{font-size:14px;color:#555}table{border-collapse:collapse;width:100%}td{border:1px solid #ccc;padding:8px}p{white-space:pre-wrap}</style></head><body>'
    output += "".join(parts) + "</body></html>"
    target.write_text(output, encoding="utf-8")


def export_document(source_path: Path | None, document: dict, plan: dict,
                    selections: list[dict], output_dir: Path, format: str) -> Path:
    format = format.lower()
    filenames = {"docx": "illustrated.docx", "markdown": "illustrated.zip", "html": "preview.html"}
    if format not in filenames:
        raise ValueError("不支持的导出格式")
    output_dir = Path(output_dir)
    target = output_dir / filenames[format]
    source_data = None
    if source_path is not None:
        source_path = Path(source_path)
        if source_path.resolve() == target.resolve():
            raise ValueError("输出不能覆盖原稿，请使用独立导出目录")
        if not source_path.is_file() or source_path.stat().st_size > MAX_FILE_BYTES:
            raise ValueError("原稿不存在或超出大小限制")
        data = source_path.read_bytes()
        if _hash(data) != document["source_hash"]:
            raise ValueError("原稿已改变，请重新导入并确认图位")
        if document["format"] == "docx":
            parsed = _parse_docx(data, source_path.name)
            source_data = data
            imported = {p["id"]: (p["text"], p["kind"], p["source_index"]) for p in document["paragraphs"]}
            actual = {p["id"]: (p["text"], p["kind"], p["source_index"]) for p in parsed["paragraphs"]}
            if imported != actual:
                raise ValueError("文档段落锚点与原稿不一致，请重新导入")
    elif document["format"] == "docx":
        raise ValueError("DOCX 导出需要保留原始文件，不能从纯文本重建")
    rows = _validated_selections(document, plan, selections)
    output_dir.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(dir=output_dir, prefix=".export-", suffix=target.suffix, delete=False)
    temporary = Path(handle.name)
    handle.close()
    try:
        if format == "docx":
            _write_docx(source_data, document, rows, temporary)
        elif format == "markdown":
            _write_markdown(document, rows, source_data, temporary)
        else:
            _write_html(document, rows, source_data, temporary)
        if source_path is not None and _hash(source_path.read_bytes()) != document["source_hash"]:
            raise ValueError("导出期间原稿已改变，请重新确认")
        os.replace(temporary, target)
        return target
    finally:
        temporary.unlink(missing_ok=True)
