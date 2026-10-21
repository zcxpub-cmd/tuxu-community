"""Public response and embedded-image leak counterexamples, with isolated data."""
import base64
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import pytest
from PIL import Image, PngImagePlugin

from backend.app.privacy import UnsafeLegacyExport, clean_export, clean_png, private_export_copy, private_image_copy, public_error, public_task
from backend.tests.test_service import approve, director_ready, docs_double, harness, ready, wait

SECRET = 'PRIVATE_RECIPE_CANARY_7a32'


def png_bytes(*, private=True, mode='RGB'):
    image = Image.new(mode, (17, 13), 1 if mode == 'P' else (31, 85, 129))
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text('Description', 'original scientific reference metadata')
    if private:
        metadata.add_text('prompt', SECRET)
        metadata.add_text('workflow', json.dumps({'private': SECRET}))
    if mode == 'P':
        image.putpalette([0, 0, 0, 31, 85, 129] + [0] * (768 - 6))
        image.info['transparency'] = 1
    output = io.BytesIO()
    image.save(output, format='PNG', pnginfo=metadata)
    return output.getvalue()


def assert_pixels_match(before, after):
    with Image.open(io.BytesIO(before)) as old, Image.open(io.BytesIO(after)) as new:
        assert old.size == new.size
        assert old.convert('RGBA').tobytes() == new.convert('RGBA').tobytes()
        assert not new.info


@pytest.mark.parametrize('mode', ['RGB', 'P'])
def test_image_metadata_copies_preserve_pixels_and_do_not_rewrite_legacy_source(tmp_path, mode):
    original = png_bytes(mode=mode)
    source = tmp_path / 'legacy.png'
    source.write_bytes(original)
    sanitized = private_image_copy(source, tmp_path / 'cache')
    assert_pixels_match(original, sanitized.read_bytes())
    assert source.read_bytes() == original
    assert private_image_copy(source, tmp_path / 'cache') == sanitized
    assert SECRET.encode() not in sanitized.read_bytes()


@pytest.mark.parametrize('mode,color,transparency', [
    ('RGB', (31, 85, 129), (31, 85, 129)), ('L', 85, 85), ('1', 0, 0),
    ('RGBA', (31, 85, 129, 37), None),
])
def test_transparent_color_and_rgba_alpha_survive_metadata_removal(mode, color, transparency):
    image = Image.new(mode, (3, 2), color)
    info = PngImagePlugin.PngInfo()
    info.add_text('prompt', SECRET)
    output = io.BytesIO()
    options = {'transparency': transparency} if transparency is not None else {}
    image.save(output, format='PNG', pnginfo=info, **options)
    original = output.getvalue()
    assert_pixels_match(original, clean_png(original))


@pytest.mark.parametrize('suffix', ['.docx', '.zip', '.html'])
def test_export_copy_cleans_only_application_generated_images_and_preserves_source_media(tmp_path, suffix):
    generated, original = png_bytes(), png_bytes(private=False)
    hashes = {hashlib.sha256(generated).hexdigest()}
    if suffix == '.html':
        generated_uri, original_uri = (base64.b64encode(image).decode() for image in (generated, original))
        data = (f'<p>caption</p><figure><img src="data:image/png;base64,{generated_uri}" alt="新配图">'
                f'<figcaption>新图说明</figcaption></figure><figure><img src="data:image/png;base64,{original_uri}" alt="原稿图片"></figure>').encode()
    else:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('word/media/generated.png' if suffix == '.docx' else 'images/figure_001.png', generated)
            archive.writestr('word/media/source.png' if suffix == '.docx' else 'images/source_0001_01.png', original)
            archive.writestr('document.xml', '<p>source text</p>')
        data = stream.getvalue()
    source = tmp_path / ('legacy' + suffix)
    source.write_bytes(data)
    cleaned = private_export_copy(source, tmp_path / 'cache', hashes).read_bytes()
    assert source.read_bytes() == data
    if suffix == '.html':
        assert base64.b64encode(original) in cleaned
        assert base64.b64encode(generated) not in cleaned
        generated_clean = base64.b64decode(re.findall(rb'data:image/png;base64,([A-Za-z0-9+/=]+)', cleaned)[0])
    else:
        with zipfile.ZipFile(io.BytesIO(cleaned)) as archive:
            generated_path = 'word/media/generated.png' if suffix == '.docx' else 'images/figure_001.png'
            source_path = 'word/media/source.png' if suffix == '.docx' else 'images/source_0001_01.png'
            generated_clean = archive.read(generated_path)
            assert archive.read(source_path) == original
            assert archive.read('document.xml') == b'<p>source text</p>'
    assert_pixels_match(generated, generated_clean)


def source_and_generated_export(tmp_path, format, *, sanitize_generated=False):
    from docx import Document
    from backend.app.documents import export_document, parse_document
    original = png_bytes()
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('原稿正文')
    doc.add_picture(io.BytesIO(original))
    doc.save(source)
    source_bytes = source.read_bytes()
    document = parse_document(source)
    paragraph = next(p for p in document['paragraphs'] if p['text'] == '原稿正文')
    image_path = tmp_path / 'generated.png'
    image_path.write_bytes(clean_png(original) if sanitize_generated else original)
    selection = {'slot_id': 's', 'paragraph_id': paragraph['id'], 'image_path': str(image_path),
                 'caption': '新增示意图', 'alt_text': '原稿图片'}
    output = export_document(source, document, {'slots': [{'id': 's', 'paragraph_id': paragraph['id']}]},
                             [selection], tmp_path / 'out', format)
    assert source.read_bytes() == source_bytes
    return output, source, source_bytes, original


@pytest.mark.parametrize('format', ['markdown', 'html'])
def test_actual_export_same_image_bytes_preserve_source_role_and_clean_only_generated_role(tmp_path, format):
    output, source, source_bytes, original = source_and_generated_export(tmp_path, format)
    raw = output.read_bytes()
    cleaned = clean_export(raw, output.suffix, {hashlib.sha256(original).hexdigest()}, source_bytes)
    if format == 'markdown':
        with zipfile.ZipFile(io.BytesIO(cleaned)) as archive:
            source_name = next(name for name in archive.namelist() if name.startswith('images/source_'))
            assert archive.read(source_name) == original
            assert_pixels_match(original, archive.read('images/figure_001.png'))
    else:
        figures = re.findall(rb'<figure>.*?</figure>', cleaned)
        assert len(figures) == 2 and b'<figcaption>' in figures[0]
        generated, protected = [base64.b64decode(re.search(rb'base64,([A-Za-z0-9+/=]+)', figure)[1]) for figure in figures]
        assert protected == original
        assert_pixels_match(original, generated)
    assert source.read_bytes() == source_bytes and output.read_bytes() == raw


def test_actual_old_docx_shared_source_generated_part_blocks_safe_download_then_clean_insert_succeeds(tmp_path):
    from backend.app.privacy import _docx_image_uses
    output, source, source_bytes, original = source_and_generated_export(tmp_path, 'docx')
    raw = output.read_bytes()
    hashes = {hashlib.sha256(original).hexdigest()}
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert list(_docx_image_uses(archive).values()) == [2]
    with pytest.raises(UnsafeLegacyExport):
        private_export_copy(output, tmp_path / 'cache', hashes, source_bytes)
    assert source.read_bytes() == source_bytes and output.read_bytes() == raw

    fresh = tmp_path / 'fresh'
    fresh.mkdir()
    output, source, source_bytes, original = source_and_generated_export(fresh, 'docx', sanitize_generated=True)
    clean_generated = clean_png(original)
    cleaned = clean_export(output.read_bytes(), '.docx', hashes | {hashlib.sha256(clean_generated).hexdigest()}, source_bytes)
    with zipfile.ZipFile(io.BytesIO(source_bytes)) as source_archive, zipfile.ZipFile(io.BytesIO(cleaned)) as archive:
        original_part = next(name for name in source_archive.namelist() if name.startswith('word/media/'))
        assert archive.read(original_part) == original
        added = [name for name in archive.namelist() if name.startswith('word/media/') and name != original_part]
        assert len(added) == 1
        assert_pixels_match(original, archive.read(added[0]))
    assert source.read_bytes() == source_bytes


@pytest.mark.parametrize('suffix', ['.zip', '.html'])
def test_unknown_legacy_media_roles_are_rejected_instead_of_silently_rewriting_source(suffix):
    original = png_bytes()
    if suffix == '.zip':
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('images/unknown.png', original)
        data = stream.getvalue()
    else:
        data = b'<img src="data:image/png;base64,' + base64.b64encode(original) + b'">'
    with pytest.raises(UnsafeLegacyExport):
        clean_export(data, suffix, {hashlib.sha256(original).hexdigest()})


def test_ambiguous_legacy_download_returns_safe_actionable_reexport_message(harness):
    client, service, _, _ = harness
    task = ready(harness)
    original = png_bytes()
    directory = service.directory(task['id']) / 'downloads'
    directory.mkdir(exist_ok=True)
    path = directory / 'legacy.html'
    raw = b'<img src="data:image/png;base64,' + base64.b64encode(original) + b'">'
    path.write_bytes(raw)
    def update(t):
        t['assets'] = [{'id': 'old-generated', 'source': 'comfyui', 'sha256': hashlib.sha256(original).hexdigest()}]
        t['exports'] = [{'filename': 'legacy.html'}]
    service.store.mutate(task['id'], update)
    response = client.get(f'/api/exports/{task["id"]}/legacy.html')
    assert response.status_code == 409 and '重新导出' in response.text and SECRET not in response.text
    assert path.read_bytes() == raw


def test_public_projection_drops_all_private_prompt_surfaces_without_mutation():
    task = {'id': 'task', 'prompts': {'s': {'text': SECRET, 'negative': SECRET, 'regions': [{'object_id': 'tree', 'text': SECRET}]}},
        'plan': {'slots': [{'id': 's', 'description': '树木示意画面', 'director': {'objects': [{'id': 'tree'}]}}]},
        'approval': {'id': 'approval', 'prompt_hash': 'hash', 'prompt_evidence': [{'raw_response': SECRET}]},
        'jobs': [{'status': 'failed', 'message': SECRET, 'error': SECRET, 'prompts': {'s': SECRET},
                  'submissions': [{'state': 'failed', 'rejection': SECRET}]}],
        'assets': [{'id': 'asset', 'prompt': SECRET, 'negative_prompt': SECRET, 'scene_prompt': SECRET, 'background_text': SECRET,
                    'global_text': SECRET,
                    'regions': [{'object_id': 'tree', 'text': SECRET}],
                    'region_areas': [{'object_id': 'tree', 'x': 8, 'text': SECRET}]}], 'error': SECRET}
    before = json.dumps(task, sort_keys=True)
    visible = public_task(task)
    assert SECRET not in json.dumps(visible)
    assert visible['prompt_summaries']['s']['description'] == '树木示意画面'
    assert visible['approval']['prompt_hash'] == 'hash'
    assert visible['assets'][0]['region_areas'][0] == {'object_id': 'tree', 'x': 8}
    assert json.dumps(task, sort_keys=True) == before


def test_global_text_is_stripped_recursively_from_all_nested_snapshot_surfaces():
    task = {'id': 'task', 'global_text': SECRET,
            'plan': {'slots': [{'id': 's', 'description': '简单描述', 'global_text': SECRET,
                               'extra': [{'global_text': SECRET, 'safe': 'visible'}]}]},
            'approval': {'global_text': SECRET, 'snapshot': {'global_text': SECRET, 'safe': 'hash'}},
            'jobs': [{'status': 'completed', 'slots': [{'global_text': SECRET}],
                      'snapshot': {'nested': [{'global_text': SECRET}]},
                      'submissions': [{'global_text': SECRET}]}],
            'assets': [{'global_text': SECRET, 'processing': {'global_text': SECRET},
                        'region_areas': [{'global_text': SECRET, 'x': 8}]}]}
    before = json.dumps(task, sort_keys=True)
    visible = public_task(task)
    assert SECRET not in json.dumps(visible) and 'global_text' not in json.dumps(visible)
    assert visible['plan']['slots'][0]['extra'] == [{'safe': 'visible'}]
    assert visible['approval']['snapshot'] == {'safe': 'hash'}
    assert json.dumps(task, sort_keys=True) == before


@pytest.mark.parametrize('private_message', [
    '提交结果未知：' + SECRET,
    '区域控制节点：' + SECRET,
    '校验拒绝：' + SECRET,
    '正向提示词含排除指令：' + SECRET,
    '全文覆盖遗漏：' + SECRET,
    '请求超时：' + SECRET,
    SECRET,
])
def test_public_error_is_exactly_idempotent_but_never_passes_an_appended_private_canary(private_message):
    safe = public_error(private_message)
    assert public_error(safe) == safe
    appended = safe + ' ' + SECRET
    assert public_error(appended) != appended and SECRET not in public_error(appended)
    prefixed = SECRET + ' ' + safe
    assert public_error(prefixed) != prefixed and SECRET not in public_error(prefixed)


def test_bootstrap_task_approve_generation_and_422_do_not_return_private_prompt(harness):
    client, service, text, comfy = harness
    task = director_ready(harness)
    def private_prompts(slots, style, model):
        return {slot['id']: {'text': SECRET + ' A tree.', 'negative': SECRET,
            'background_text': SECRET + ' A neutral background and support surface.',
            'regions': [{'object_id': item['id'], 'text': SECRET + ' A tree.'} for item in slot['director']['objects']]}
            for slot in slots}
    text.prompts = private_prompts
    task = approve(client, task)
    assert 'prompts' not in task and SECRET not in json.dumps(task)
    internal = service.require(task['id'])
    assert SECRET in internal['prompts'][task['plan']['slots'][0]['id']]['text']
    assert SECRET in internal['prompts'][task['plan']['slots'][0]['id']]['global_text']
    response = client.post(f'/api/tasks/{task["id"]}/generate', json={
        'version': task['plan']['version'], 'approval_id': task['approval']['id'], 'idempotency_key': 'privacy'})
    assert SECRET not in response.text
    task, job = wait(client, task['id'], response.json()['job_id'])
    for url in ('/api/bootstrap', '/api/tasks', '/api/tasks/' + task['id']):
        assert SECRET not in client.get(url).text
    # Default FastAPI validation errors echo input; the public handler must not.
    response = client.post('/api/tasks/paste', json=[{'text': SECRET}])
    assert response.status_code == 422 and SECRET not in response.text
    response = client.put(f'/api/tasks/{task["id"]}/prompts', json={'prompts': {SECRET: SECRET}})
    assert response.status_code == 410 and SECRET not in response.text


def test_legacy_asset_and_every_processed_version_serve_metadata_free_pixels(harness):
    client, service, _, _ = harness
    task = approve(client, ready(harness))
    response = client.post(f'/api/tasks/{task["id"]}/generate', json={
        'version': task['plan']['version'], 'approval_id': task['approval']['id'], 'idempotency_key': 'legacy-private-image'})
    task, _ = wait(client, task['id'], response.json()['job_id'])
    asset = task['assets'][0]
    legacy = png_bytes()
    asset_path = service.directory(task['id']) / 'assets' / (asset['id'] + '.png')
    asset_path.write_bytes(legacy)
    for version in ('v-one', 'v-two'):
        asset_path.with_name(asset['id'] + '-' + version + '.png').write_bytes(legacy)
    def update(t):
        stored = t['assets'][0]
        stored['sha256'] = hashlib.sha256(legacy).hexdigest()
        stored['processing_versions'] = {version: {'filename': asset['id'] + '-' + version + '.png',
            'processing': {'sha256': hashlib.sha256(legacy).hexdigest()}} for version in ('v-one', 'v-two')}
        stored['processed_url'] = f'/api/assets/{task["id"]}/{asset["id"]}/processed?version=v-two'
        stored['processing'] = {'version': 'v-two'}
    service.store.mutate(task['id'], update)
    before = service.require(task['id'])
    for url in [asset['url']] + [f'/api/assets/{task["id"]}/{asset["id"]}/processed?version={version}' for version in ('v-one', 'v-two')]:
        response = client.get(url)
        assert response.status_code == 200
        assert response.headers['X-Content-SHA256'] == hashlib.sha256(response.content).hexdigest()
        assert_pixels_match(legacy, response.content)
    assert asset_path.read_bytes() == legacy and service.require(task['id']) == before
